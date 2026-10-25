"""Lectores de Word y Excel y entrega al mismo recorrido de revisión; sin API real."""
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED
import json
import unittest

import httpx
from docx import Document
from google import genai
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest
import app
from test_extraccion import datos_extraidos, imagen
from test_ia import CLAVE_FICTICIA


def word(imagen_incrustada=False):
    doc = Document()
    doc.add_paragraph("Proveedor ficticio; moneda USD")
    tabla = doc.add_table(rows=2, cols=5)
    for fila, valores in zip(tabla.rows, (("Código", "Cantidad", "UOM", "Precio USD", "MOQ"),
                                         ("VAL-100", "100", "piezas", "10.00", "50"))):
        for celda, valor in zip(fila.cells, valores):
            celda.text = valor
    doc.add_paragraph("DDP Incoterms 2020; Almacén de Ejemplo; 30 días calendario hasta recepción.")
    doc.sections[0].header.paragraphs[0].text = "Cotización SQ-DEMO-001"
    doc.sections[0].footer.paragraphs[0].text = "Pago a 30 días desde factura"
    if imagen_incrustada:
        doc.add_picture(BytesIO(imagen()))
    salida = BytesIO()
    doc.save(salida)
    return salida.getvalue()


def excel(formula=False, oculta=False, vacio=False):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Oferta"
    if not vacio:
        for fila in (("Proveedor ficticio", "USD"), ("Código", "Cantidad", "Precio unitario", "MOQ"),
                     ("VAL-100", 100, "=5+5" if formula else 10, 50), ("Gasto adicional confirmado", 0)):
            hoja.append(fila)
        hoja["C3"].number_format = '0.00 "USD"'
        hoja["C4"] = "#DIV/0!" if formula else None
    if oculta:
        otra = libro.create_sheet("Condiciones")
        otra["A1"] = "MOQ 50; sin múltiplo obligatorio"
        otra.sheet_state = "hidden"
    salida = BytesIO()
    libro.save(salida)
    libro.close()
    return salida.getvalue()


def reemplazar_parte(contenido, nombre, modificar):
    salida = BytesIO()
    with ZipFile(BytesIO(contenido)) as original, ZipFile(salida, "w", ZIP_DEFLATED) as nuevo:
        for entrada in original.infolist():
            datos = original.read(entrada)
            nuevo.writestr(entrada.filename, modificar(datos) if entrada.filename == nombre else datos)
    return salida.getvalue()


class LecturaDeOficina(unittest.TestCase):
    def test_word_conserva_parrafos_tablas_encabezado_y_pie_con_referencias(self):
        fuente = app.validar_archivo_cotizacion("../oferta.DOCX", word())
        self.assertEqual(fuente["archivo"], "oferta.DOCX")
        self.assertIsNone(fuente["paginas"])
        registros = json.loads(fuente["texto_preparado"])["registros"]
        self.assertTrue(any("tabla" in r["origen"] and "VAL-100" in r["texto"] for r in registros))
        self.assertTrue(any("header" in r["origen"] and "SQ-DEMO-001" in r["texto"] for r in registros))
        self.assertTrue(any("footer" in r["origen"] and "30 días" in r["texto"] for r in registros))
        self.assertIn("Incoterms 2020", fuente["texto_preparado"])

    def test_excel_conserva_hojas_coordenadas_formatos_ceros_y_vacios(self):
        fuente = app.validar_archivo_cotizacion("oferta.xlsx", excel(oculta=True))
        registros = json.loads(fuente["texto_preparado"])["registros"]
        precio = next(r for r in registros if r["origen"].endswith("'Oferta', celda C3"))
        self.assertEqual(precio["valor"], 10)
        self.assertEqual(precio["formato"], '0.00 "USD"')
        self.assertEqual(next(r for r in registros if r["origen"].endswith("'Oferta', celda B4"))["valor"], 0)
        self.assertFalse(any(r["origen"].endswith("'Oferta', celda C4") for r in registros))
        self.assertTrue(any(r["hoja_oculta"] and "MOQ 50" in r.get("valor", "") for r in registros))
        self.assertTrue(any("oculta" in a for a in fuente["avisos_archivo"]))

    def test_formulas_y_errores_no_se_presentan_como_valores_calculados(self):
        fuente = app.validar_archivo_cotizacion("oferta.xlsx", excel(formula=True))
        registros = json.loads(fuente["texto_preparado"])["registros"]
        formula = next(r for r in registros if r["origen"].endswith("'Oferta', celda C3"))
        self.assertEqual(formula["formula_sin_resultado_confirmado"], "=5+5")
        self.assertNotIn("valor", formula)
        self.assertTrue(any(r.get("error_celda") == "#DIV/0!" for r in registros))
        self.assertTrue(any("fórmulas" in a for a in fuente["avisos_archivo"]))

    def test_imagenes_incrustadas_tienen_aviso_sin_ocultar_el_texto(self):
        fuente = app.validar_archivo_cotizacion("oferta.docx", word(imagen_incrustada=True))
        self.assertIn("VAL-100", fuente["texto_preparado"])
        self.assertTrue(any("imágenes" in a for a in fuente["avisos_archivo"]))

    def test_formato_real_incorrecto_archivo_danado_y_versiones_antiguas_se_rechazan(self):
        for nombre, contenido in (("oferta.docx", excel()), ("oferta.xlsx", word()),
                                  ("oferta.docx", b"PK-roto"), ("oferta.xlsx", b"\xd0\xcf\x11\xe0"),
                                  ("oferta.doc", word()), ("oferta.xls", excel())):
            with self.subTest(nombre=nombre), self.assertRaises(ValueError):
                app.validar_archivo_cotizacion(nombre, contenido)

    def test_excel_admite_tipo_ooxml_por_defecto_y_prioriza_el_tipo_explicito(self):
        from xml.etree import ElementTree
        tipo_excel = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
        def tipo_por_defecto(dato):
            raiz = ElementTree.fromstring(dato)
            for nodo in list(raiz):
                if nodo.attrib.get("PartName") == "/xl/workbook.xml":
                    raiz.remove(nodo)
                elif nodo.attrib.get("Extension") == "xml":
                    nodo.set("ContentType", tipo_excel)
            return ElementTree.tostring(raiz, encoding="utf-8", xml_declaration=True)
        contenido = reemplazar_parte(excel(), "[Content_Types].xml", tipo_por_defecto)
        self.assertIn("VAL-100", app.validar_archivo_cotizacion("oferta.xlsx", contenido)["texto_preparado"])
        def macro_explicita(dato):
            raiz = ElementTree.fromstring(dato)
            ElementTree.SubElement(raiz, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", PartName="/xl/workbook.xml", ContentType="application/vnd.ms-excel.sheet.macroEnabled.main+xml")
            return ElementTree.tostring(raiz, encoding="utf-8", xml_declaration=True)
        contenido_macro = reemplazar_parte(contenido, "[Content_Types].xml", macro_explicita)
        with self.assertRaises(ValueError):
            app.validar_archivo_cotizacion("oferta.xlsx", contenido_macro)

    def test_archivo_vacio_no_envia_una_lectura_aparentemente_completa(self):
        vacio_word = BytesIO()
        Document().save(vacio_word)
        for nombre, contenido in (("oferta.docx", vacio_word.getvalue()), ("oferta.xlsx", excel(vacio=True))):
            with self.subTest(nombre=nombre), self.assertRaisesRegex(ValueError, "No se encontró"):
                app.validar_archivo_cotizacion(nombre, contenido)

    def test_word_con_cuadros_de_texto_o_cambios_requiere_pdf(self):
        for etiqueta in ("txbxContent", "ins", "del", "altChunk"):
            contenido = reemplazar_parte(word(), "word/document.xml", lambda dato: dato.replace(b"<w:body>", b"<w:body><w:" + etiqueta.encode() + b"/>"))
            with self.subTest(etiqueta=etiqueta), self.assertRaisesRegex(ValueError, "PDF"):
                app.validar_archivo_cotizacion("oferta.docx", contenido)

    def test_entidades_xml_y_macros_no_se_procesan(self):
        malicioso = reemplazar_parte(word(), "word/document.xml", lambda _: b'<!DOCTYPE x [<!ENTITY secreto SYSTEM "file:///ejemplo-inexistente">]><x>&secreto;</x>')
        macro = reemplazar_parte(excel(), "[Content_Types].xml", lambda dato: dato.replace(b"spreadsheetml.sheet.main+xml", b"ms-excel.sheet.macroEnabled.main+xml"))
        for nombre, contenido in (("oferta.docx", malicioso), ("oferta.xlsx", macro)):
            with self.subTest(nombre=nombre), self.assertRaises(ValueError):
                app.validar_archivo_cotizacion(nombre, contenido)

    def test_limites_de_hojas_filas_columnas_celdas_y_texto_fallan_antes_de_la_api(self):
        for limite, valor, contenido in (("MAX_HOJAS_EXCEL", 1, excel(oculta=True)),
                                         ("MAX_FILAS_EXCEL", 2, excel()), ("MAX_COLUMNAS_EXCEL", 2, excel()),
                                         ("MAX_CELDAS_EXCEL", 2, excel()), ("MAX_TEXTO_OFICINA", 100, excel())):
            with self.subTest(limite=limite), patch.object(app, limite, valor), patch("google.genai.Client") as cliente:
                resultado = app.extraer_archivo_ia("oferta.xlsx", contenido)
                self.assertFalse(resultado["ok"])
                self.assertIn("supera", resultado["mensaje"])
                cliente.assert_not_called()
        with patch.object(app, "MAX_OFICINA_DESCOMPRIMIDO", 100), self.assertRaisesRegex(ValueError, "abrirlo"):
            app.validar_archivo_cotizacion("oferta.docx", word())

    def test_dimensiones_de_excel_incorrectas_no_ocultan_celdas(self):
        contenido = reemplazar_parte(excel(), "xl/worksheets/sheet1.xml", lambda dato: __import__("re").sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', dato))
        fuente = app.validar_archivo_cotizacion("oferta.xlsx", contenido)
        self.assertIn("celda C3", fuente["texto_preparado"])

    def test_excepcion_del_lector_no_revela_documentos_ni_detalles_internos(self):
        for extension, lector, contenido in (("docx", "leer_word", word()), ("xlsx", "leer_excel", excel())):
            with self.subTest(extension=extension), patch.object(app, lector, side_effect=ValueError("DETALLE_PRIVADO_FICTICIO")):
                resultado = app.extraer_archivo_ia("oferta." + extension, contenido)
                self.assertFalse(resultado["ok"])
                self.assertIn("manualmente", resultado["mensaje"])
                self.assertNotIn("DETALLE_PRIVADO_FICTICIO", resultado["mensaje"])


class ExtraccionYRevisionDeOficina(unittest.TestCase):
    def test_se_envia_texto_con_referencias_sin_binario_y_sin_relajar_el_contrato(self):
        for nombre, contenido in (("oferta.docx", word()), ("oferta.xlsx", excel(formula=True))):
            with self.subTest(nombre=nombre), patch.object(app.st, "secrets", {"GEMINI_API_KEY": CLAVE_FICTICIA}), patch("google.genai.Client") as crear:
                cliente = crear.return_value.__enter__.return_value
                cliente.models.generate_content.return_value = SimpleNamespace(text=json.dumps(datos_extraidos()), candidates=[])
                resultado = app.extraer_archivo_ia(nombre, contenido)
                self.assertTrue(resultado["ok"], resultado)
                peticion = cliente.models.generate_content.call_args.kwargs
                self.assertIsInstance(peticion["contents"][1], str)
                self.assertIn("VAL-100", peticion["contents"][1])
                self.assertIn("registros", peticion["contents"][1])
                self.assertNotIn(CLAVE_FICTICIA, str(peticion))
                self.assertIn("No las calcules", peticion["config"].system_instruction)
                self.assertFalse(peticion["config"].response_json_schema["additionalProperties"])
                self.assertNotIn("texto_preparado", resultado["campos"]["extraccion"])
                if nombre.endswith("xlsx"):
                    self.assertTrue(any("fórmulas" in a for a in resultado["campos"]["extraccion"]["avisos"]))
                cliente.files.upload.assert_not_called()

    def test_sdk_real_serializa_texto_de_word_y_excel_sin_subir_archivos(self):
        for nombre, contenido in (("oferta.docx", word()), ("oferta.xlsx", excel())):
            solicitudes = []
            def responder(request):
                solicitudes.append(json.loads(request.content))
                return httpx.Response(200, json={"candidates": [{"content": {"role": "model", "parts": [{"text": json.dumps(datos_extraidos())}]}, "finishReason": "STOP"}]})
            cliente_real = genai.Client
            def crear_simulado(**kwargs):
                kwargs["http_options"].httpx_client = httpx.Client(transport=httpx.MockTransport(responder))
                return cliente_real(**kwargs)
            with self.subTest(nombre=nombre), patch.object(app.st, "secrets", {"GEMINI_API_KEY": CLAVE_FICTICIA}), patch("google.genai.Client", side_effect=crear_simulado):
                self.assertTrue(app.extraer_archivo_ia(nombre, contenido)["ok"])
            self.assertEqual(len(solicitudes), 1)
            partes = solicitudes[0]["contents"][0]["parts"]
            self.assertTrue(all("text" in parte for parte in partes))
            self.assertTrue(any("VAL-100" in parte["text"] for parte in partes))

    def test_archivos_incorrectos_preservan_entrada_manual_y_no_llaman_a_ia(self):
        for nombre in ("oferta.docx", "oferta.xlsx"):
            with self.subTest(nombre=nombre), patch("google.genai.Client") as cliente:
                resultado = app.extraer_archivo_ia(nombre, b"archivo ficticio roto")
                self.assertFalse(resultado["ok"])
                self.assertIn("manual", resultado["mensaje"])
                self.assertEqual(app.total_original(app.crear_cotizacion({"proveedor": "Prueba", "producto": "Válvula", "cantidad": "2", "precio_unitario": "10", "envio": "0", "impuestos": "0", "moneda": "USD"})), 20)
                cliente.assert_not_called()

    def test_word_llega_a_revision_sin_confirmarse_automaticamente(self):
        archivo = SimpleNamespace(name="oferta.docx", getvalue=word)
        with patch("streamlit.file_uploader", return_value=archivo), patch("google.genai.Client") as crear:
            cliente = crear.return_value.__enter__.return_value
            cliente.models.generate_content.return_value = SimpleNamespace(text=json.dumps(datos_extraidos()), candidates=[])
            at = AppTest.from_file(str(Path(app.__file__).resolve()), default_timeout=30)
            at.secrets["GEMINI_API_KEY"] = CLAVE_FICTICIA
            at.run().button("agregar").click().run()
            at.checkbox("aceptar_carga_0").check().run()
            at.button("extraer_archivo_0").click().run()
            self.assertFalse(at.exception)
            self.assertEqual(len(at.session_state["cotizaciones_confirmadas"]), 0)
            self.assertEqual(at.session_state["campos_formulario"]["extraccion"]["archivo"], "oferta.docx")
            next(b for b in at.button if b.label == "Revisar cotización").click().run()
            self.assertFalse(at.exception)
            self.assertIsNotNone(at.session_state["borrador"])
            self.assertEqual(len(at.session_state["cotizaciones_confirmadas"]), 0)


if __name__ == "__main__":
    unittest.main()
